"""Document processing: validation and text extraction for PDF, DOCX and TXT.

    uploaded bytes -> validate (type, size, magic bytes) -> parse -> raw text
Cleaning/NLP happens later in preprocessing.py.
"""
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePath

from app.config import get_settings
from app.utils.errors import ExtractionFailed, FileTooLarge, InsufficientText, UnsupportedFileType

ALLOWED_EXTENSIONS = {"pdf", "docx", "txt"}


@dataclass
class ParsedDocument:
    filename: str
    file_type: str
    size: int
    text: str


def safe_filename(name: str | None) -> str:
    """Strip any path components and unusual characters from a client-supplied name."""
    base = PurePath((name or "document").replace("\\", "/")).name
    base = re.sub(r"[^\w.\- ()]", "_", base).strip(" .") or "document"
    return base[:150]


def get_extension(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def validate_upload(filename: str, content: bytes) -> str:
    """Check extension, size and file signature. Returns the normalised extension."""
    settings = get_settings()
    ext = get_extension(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise UnsupportedFileType(ext)
    if len(content) == 0:
        raise ExtractionFailed("This file is empty.", "Choose a file that contains your resume.")
    if len(content) > settings.max_upload_bytes:
        raise FileTooLarge(settings.max_upload_mb)
    if ext == "pdf" and not content.lstrip()[:5].startswith(b"%PDF"):
        raise ExtractionFailed("This file has a .pdf extension but is not a valid PDF.")
    if ext == "docx" and not content[:2] == b"PK":
        raise ExtractionFailed("This file has a .docx extension but is not a valid Word document.")
    return ext


def _parse_pdf(content: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            try:
                if reader.decrypt("") == 0:
                    raise ExtractionFailed("This PDF is password-protected.",
                                           "Remove the password and upload it again.")
            except ExtractionFailed:
                raise
            except Exception:
                raise ExtractionFailed("This PDF is password-protected.",
                                       "Remove the password and upload it again.")
        pages = []
        for page in reader.pages[:30]:
            pages.append(page.extract_text() or "")
        return "\n".join(pages)
    except ExtractionFailed:
        raise
    except (PyPdfError, ValueError, KeyError, OSError, TypeError):
        raise ExtractionFailed("We couldn't read this PDF. It may be corrupted.")
    except Exception:
        raise ExtractionFailed("We couldn't read this PDF.")


def _parse_docx(content: bytes) -> str:
    from docx import Document

    try:
        if not zipfile.is_zipfile(io.BytesIO(content)):
            raise ExtractionFailed("This Word document appears to be corrupted.")
        doc = Document(io.BytesIO(content))
        parts: list[str] = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                parts.append("")
                continue
            style = (para.style.name or "").lower() if para.style is not None else ""
            parts.append(f"- {text}" if "list" in style else text)
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(dict.fromkeys(cells)))
        return "\n".join(parts)
    except ExtractionFailed:
        raise
    except Exception:
        raise ExtractionFailed("We couldn't read this Word document. It may be corrupted.")


def _parse_txt(content: bytes) -> str:
    if b"\x00" in content[:4096]:
        raise ExtractionFailed("This file doesn't look like a text file.")
    encodings = ["utf-8-sig"]
    if content[:2] in (b"\xff\xfe", b"\xfe\xff"):  # UTF-16 only when a BOM proves it (else it decodes any bytes)
        encodings.append("utf-16")
    encodings += ["cp1252", "latin-1"]
    for enc in encodings:
        try:
            return content.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise ExtractionFailed("We couldn't decode this text file.")


def extract_text(filename: str, content: bytes, what: str = "resume", min_chars: int | None = None,
                 min_words: int | None = None) -> ParsedDocument:
    """Validate and parse an uploaded file into plain text."""
    settings = get_settings()
    filename = safe_filename(filename)
    ext = validate_upload(filename, content)
    parser = {"pdf": _parse_pdf, "docx": _parse_docx, "txt": _parse_txt}[ext]
    text = parser(content)[: settings.max_text_chars]
    text = text.replace("\x00", "")
    check_sufficient(text, what, min_chars, min_words)
    return ParsedDocument(filename=filename, file_type=ext, size=len(content), text=text)


def check_sufficient(text: str, what: str = "resume", min_chars: int | None = None,
                     min_words: int | None = None) -> None:
    settings = get_settings()
    min_chars = settings.min_resume_chars if min_chars is None else min_chars
    min_words = settings.min_resume_words if min_words is None else min_words
    stripped = text.strip()
    if len(stripped) < min_chars or len(stripped.split()) < min_words:
        raise InsufficientText(what)
