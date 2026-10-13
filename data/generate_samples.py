"""Generate sample_resume.pdf and sample_resume.docx from sample_resume.txt."""
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

HERE = Path(__file__).parent
HEADINGS = {"SUMMARY", "EDUCATION", "TECHNICAL SKILLS", "EXPERIENCE", "PROJECTS",
            "CERTIFICATIONS", "ACHIEVEMENTS", "LANGUAGES"}
lines = (HERE / "sample_resume.txt").read_text(encoding="utf-8").splitlines()

# ---- PDF
styles = getSampleStyleSheet()
body = ParagraphStyle("b", parent=styles["BodyText"], fontSize=9.5, leading=12.5)
head = ParagraphStyle("h", parent=styles["Heading3"], fontSize=11, spaceBefore=8, spaceAfter=2)
name = ParagraphStyle("n", parent=styles["Title"], fontSize=18, alignment=0)
story = []
for i, line in enumerate(lines):
    esc = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    if not line.strip():
        story.append(Spacer(1, 3))
    elif i == 0:
        story.append(Paragraph(esc, name))
    elif line in HEADINGS:
        story.append(Paragraph(esc, head))
    elif line.startswith("- "):
        story.append(Paragraph("&bull; " + esc[2:], ParagraphStyle("li", parent=body, leftIndent=10, firstLineIndent=-8)))
    else:
        story.append(Paragraph(esc, body))
SimpleDocTemplate(str(HERE / "sample_resume.pdf"), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                  topMargin=15 * mm, bottomMargin=15 * mm).build(story)

# ---- DOCX
doc = Document()
for i, line in enumerate(lines):
    if not line.strip():
        continue
    if i == 0:
        doc.add_heading(line, level=0)
    elif line in HEADINGS:
        doc.add_heading(line, level=2)
    elif line.startswith("- "):
        doc.add_paragraph(line[2:], style="List Bullet")
    else:
        doc.add_paragraph(line)
doc.save(str(HERE / "sample_resume.docx"))
print("generated pdf + docx")
